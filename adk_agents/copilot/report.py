"""基于查询快照生成三部分分析报告，并导出 Word。"""

from __future__ import annotations

import json
import re
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any
from uuid import uuid4

from adk_agents.copilot.chart import (
    coerce_chart_plan,
    format_snapshot_result_overview,
    infer_fallback_chart_plan,
    materialize_chart_spec,
    parse_json_objects,
)

MAX_REPORT_ANALYSES = 3
REPORTS_DIR = Path(__file__).resolve().parents[2] / "data" / "reports"

_UNSAFE_FILENAME_RE = re.compile(r'[\\/:*?"<>|\s]+')


def reports_dir() -> Path:
    path = REPORTS_DIR
    path.mkdir(parents=True, exist_ok=True)
    return path


def _empty_report(summary: str, *, source_row_count: int = 0) -> dict[str, Any]:
    return {
        "kind": "analysis_report",
        "feasible": False,
        "summary": summary,
        "title": "",
        "overview": "",
        "conclusions": "",
        "analyses": [],
        "file_id": "",
        "filename": "",
        "source_row_count": source_row_count,
    }


def format_report_planner_user_message(
        *,
        user_request: str,
        snapshot: dict[str, Any],
) -> str:
    previous_question = str(snapshot.get("user_question") or "").strip() or "（无）"
    tables = snapshot.get("selected_tables")
    table_lines = (
        "\n".join(f"- `{name}`" for name in tables if str(name).strip())
        if isinstance(tables, list) and tables
        else "（无）"
    )
    schema_context = str(snapshot.get("schema_context") or "").strip() or "（无）"
    return "\n".join([
        "## 用户报告要求",
        "",
        (user_request or "").strip() or "请基于当前查询结果生成分析报告。",
        "",
        "## 上一轮用户问题",
        "",
        previous_question,
        "",
        "## 已选数据表",
        "",
        table_lines,
        "",
        "## 检索与表结构上下文",
        "",
        schema_context,
        "",
        "## 当前查询结果概要",
        "",
        format_snapshot_result_overview(snapshot),
        "",
        "请只输出 JSON。列名必须来自上述结果概要。正文用中文，结论要能被当前数字支撑。",
    ])


def _paragraphs(text: str) -> list[str]:
    chunks = [item.strip() for item in re.split(r"\n{2,}", (text or "").strip()) if item.strip()]
    if chunks:
        return chunks
    lines = [item.strip() for item in (text or "").splitlines() if item.strip()]
    return lines


def _conclusions_text(raw: object) -> str:
    if isinstance(raw, list):
        items = [str(item).strip() for item in raw if str(item).strip()]
        if not items:
            return ""
        numbered = []
        for index, item in enumerate(items, start=1):
            if re.match(r"^\d+[\.、]", item):
                numbered.append(item)
            else:
                numbered.append(f"{index}、{item}")
        return "\n".join(numbered)
    return str(raw or "").strip()


def _sanitize_filename(title: str) -> str:
    stem = _UNSAFE_FILENAME_RE.sub("_", (title or "").strip())[:40].strip("._")
    if not stem:
        stem = "分析报告"
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    return f"{stem}_{stamp}.docx"


def _configure_matplotlib() -> Any:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    plt.rcParams["axes.unicode_minus"] = False
    candidates = [
        "Microsoft YaHei",
        "SimHei",
        "PingFang SC",
        "Noto Sans CJK SC",
        "Source Han Sans SC",
    ]
    try:
        from matplotlib import font_manager

        available = {item.name for item in font_manager.fontManager.ttflist}
        chosen = next((name for name in candidates if name in available), None)
        if chosen:
            plt.rcParams["font.sans-serif"] = [chosen, "DejaVu Sans"]
        else:
            plt.rcParams["font.sans-serif"] = candidates + ["DejaVu Sans"]
    except Exception:
        plt.rcParams["font.sans-serif"] = candidates + ["DejaVu Sans"]
    return plt


def render_chart_png(chart: dict[str, Any]) -> bytes:
    plt = _configure_matplotlib()
    categories = [str(item) for item in (chart.get("categories") or [])]
    series = chart.get("series") if isinstance(chart.get("series"), list) else []
    palette = [str(item) for item in (chart.get("palette") or []) if str(item).strip()]
    chart_type = str(chart.get("type") or "bar").strip().lower()
    title = str(chart.get("title") or "").strip()

    fig, ax = plt.subplots(figsize=(8.2, 4.4), dpi=140)
    fig.patch.set_facecolor("#ffffff")
    ax.set_facecolor("#ffffff")
    if title:
        ax.set_title(title, fontsize=12, pad=10)

    if chart_type == "pie" and series:
        values = [float(item) for item in (series[0].get("values") or [])][: len(categories)]
        if not values or not any(value > 0 for value in values):
            ax.text(0.5, 0.5, "暂无数据", ha="center", va="center")
        else:
            colors = palette[: len(values)] or None
            ax.pie(values, labels=categories[: len(values)], colors=colors, autopct="%1.1f%%")
            ax.axis("equal")
    else:
        x_pos = list(range(len(categories)))
        if chart_type == "line":
            for index, item in enumerate(series):
                if not isinstance(item, dict):
                    continue
                values = [float(value) for value in (item.get("values") or [])]
                color = palette[index] if index < len(palette) else None
                ax.plot(
                    x_pos[: len(values)],
                    values,
                    marker="o",
                    label=str(item.get("name") or f"系列{index + 1}"),
                    color=color,
                )
        else:
            width = 0.72 / max(1, len(series))
            for index, item in enumerate(series):
                if not isinstance(item, dict):
                    continue
                values = [float(value) for value in (item.get("values") or [])]
                offset = (index - (len(series) - 1) / 2) * width
                color = palette[index] if index < len(palette) else None
                ax.bar(
                    [pos + offset for pos in x_pos[: len(values)]],
                    values,
                    width=width,
                    label=str(item.get("name") or f"系列{index + 1}"),
                    color=color,
                )
        if categories:
            ax.set_xticks(x_pos)
            ax.set_xticklabels(categories, rotation=25, ha="right")
        if len(series) > 1:
            ax.legend()
        ax.spines["top"].set_visible(False)
        ax.spines["right"].set_visible(False)
        ax.yaxis.grid(True, linestyle="--", alpha=0.35)
        ax.set_axisbelow(True)

    fig.tight_layout()
    buffer = BytesIO()
    fig.savefig(buffer, format="png", bbox_inches="tight")
    plt.close(fig)
    return buffer.getvalue()


def _set_run_font(run, *, size_pt: float = 11, bold: bool = False) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run.bold = bold
    run.font.size = Pt(size_pt)
    run.font.name = "Microsoft YaHei"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def write_report_docx(payload: dict[str, Any], destination: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt
    from docx.oxml.ns import qn

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    style = document.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(11)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    title = str(payload.get("title") or "数据分析报告").strip()
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(title)
    _set_run_font(run, size_pt=18, bold=True)

    part_one = document.add_paragraph()
    run = part_one.add_run("一、总体情况")
    _set_run_font(run, size_pt=14, bold=True)
    for block in _paragraphs(str(payload.get("overview") or "")):
        para = document.add_paragraph()
        run = para.add_run(block)
        _set_run_font(run)

    part_two = document.add_paragraph()
    run = part_two.add_run("二、分类分析")
    _set_run_font(run, size_pt=14, bold=True)
    analyses = payload.get("analyses")
    if isinstance(analyses, list):
        for index, item in enumerate(analyses, start=1):
            if not isinstance(item, dict):
                continue
            heading_text = str(item.get("heading") or f"分析维度 {index}").strip()
            sub = document.add_paragraph()
            run = sub.add_run(f"{index}、{heading_text}")
            _set_run_font(run, size_pt=12, bold=True)
            for block in _paragraphs(str(item.get("narrative") or "")):
                para = document.add_paragraph()
                run = para.add_run(block)
                _set_run_font(run)
            chart = item.get("chart")
            if isinstance(chart, dict) and chart.get("categories"):
                image_bytes = render_chart_png(chart)
                image_para = document.add_paragraph()
                image_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = image_para.add_run()
                run.add_picture(BytesIO(image_bytes), width=Inches(6.1))
                caption = str(chart.get("title") or heading_text).strip()
                cap = document.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = cap.add_run(caption)
                _set_run_font(run, size_pt=9)
            elif str(item.get("chart_error") or "").strip():
                para = document.add_paragraph()
                run = para.add_run(str(item.get("chart_error")).strip())
                _set_run_font(run, size_pt=9)

    part_three = document.add_paragraph()
    run = part_three.add_run("三、分析结论")
    _set_run_font(run, size_pt=14, bold=True)
    for block in _paragraphs(_conclusions_text(payload.get("conclusions"))):
        para = document.add_paragraph()
        run = para.add_run(block)
        _set_run_font(run)

    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(destination))


def _materialize_one_chart(
        chart_plan: dict[str, Any] | None,
        snapshot: dict[str, Any] | None,
        *,
        heading: str = "",
        user_request: str = "",
        palette: object = None,
) -> tuple[dict[str, Any] | None, str]:
    """按单个分析维度绑图；失败时按标题和结果列补一张。"""
    planned = coerce_chart_plan(chart_plan, title=heading) if chart_plan else None
    if planned and heading and not str(planned.get("title") or "").strip():
        planned = {**planned, "title": heading}

    def _try(plan: dict[str, Any] | None) -> tuple[dict[str, Any] | None, str]:
        if not plan:
            return None, ""
        spec = materialize_chart_spec(
            json.dumps(
                {"feasible": True, "palette": palette, "charts": [plan]},
                ensure_ascii=False,
            ),
            snapshot,
            user_request=user_request,
        )
        charts = spec.get("charts") if isinstance(spec.get("charts"), list) else []
        chart = charts[0] if charts and isinstance(charts[0], dict) else None
        if chart and chart.get("categories"):
            return chart, ""
        return None, str(spec.get("summary") or "").strip()

    chart, error = _try(planned)
    if chart:
        return chart, ""
    fallback = infer_fallback_chart_plan(heading, snapshot)
    if fallback:
        chart, fallback_error = _try(fallback)
        if chart:
            return chart, ""
        error = error or fallback_error
    return None, error or "当前结果无法按该维度绘图。"


def materialize_analysis_report(
        raw_text: str,
        snapshot: dict[str, Any] | None,
        *,
        user_request: str = "",
        write_file: bool = True,
) -> dict[str, Any]:
    source_row_count = 0
    if snapshot:
        source_row_count = int(snapshot.get("row_count") or 0)
    parsed_objects = parse_json_objects(raw_text)
    plan = parsed_objects[-1] if parsed_objects else {}
    if not plan:
        return _empty_report(
            "未能理解报告方案，请再说明希望分析哪些维度。",
            source_row_count=source_row_count,
        )
    if plan.get("feasible") is False:
        return _empty_report(
            str(plan.get("summary") or "当前结果不足以生成分析报告。").strip(),
            source_row_count=source_row_count,
        )

    title = str(plan.get("title") or "").strip() or str(
        (snapshot or {}).get("user_question") or "数据分析报告"
    ).strip()
    overview = str(plan.get("overview") or "").strip()
    conclusions = _conclusions_text(plan.get("conclusions"))
    raw_analyses = plan.get("analyses")
    if not isinstance(raw_analyses, list):
        raw_analyses = []

    analyses: list[dict[str, Any]] = []
    for item in raw_analyses[:MAX_REPORT_ANALYSES]:
        if not isinstance(item, dict):
            continue
        heading = str(item.get("heading") or item.get("dimension") or "").strip() or "分析维度"
        narrative = str(item.get("narrative") or item.get("analysis") or "").strip()
        chart_plan = item.get("chart") if isinstance(item.get("chart"), dict) else None
        chart, chart_error = _materialize_one_chart(
            chart_plan,
            snapshot,
            heading=heading,
            user_request=user_request,
            palette=plan.get("palette"),
        )
        analyses.append({
            "heading": heading,
            "narrative": narrative,
            "chart": chart,
            "chart_error": "" if chart else chart_error,
        })

    if not overview and not conclusions and not analyses:
        return _empty_report(
            "报告正文为空，请再试一次。",
            source_row_count=source_row_count,
        )

    payload: dict[str, Any] = {
        "kind": "analysis_report",
        "feasible": True,
        "summary": str(plan.get("summary") or "已根据当前查询结果生成分析报告。").strip(),
        "title": title,
        "overview": overview,
        "conclusions": conclusions,
        "analyses": analyses,
        "file_id": "",
        "filename": "",
        "source_row_count": source_row_count,
    }
    if write_file:
        file_id = uuid4().hex
        filename = _sanitize_filename(title)
        path = reports_dir() / f"{file_id}.docx"
        write_report_docx(payload, path)
        payload["file_id"] = file_id
        payload["filename"] = filename
    return payload


def resolve_report_file(file_id: str) -> Path | None:
    token = re.sub(r"[^a-fA-F0-9]", "", file_id or "")
    if len(token) != 32:
        return None
    path = reports_dir() / f"{token}.docx"
    if not path.is_file():
        return None
    return path
